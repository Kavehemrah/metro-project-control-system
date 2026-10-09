from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import jdatetime
import openpyxl

from app.db import connect
from app.services.control import activity_control, detect_finance_discrepancy, scan_formula_errors
from app.services.forecast import calculate_schedule


def _clean_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\u200c", " ").replace("\n", " ").strip()


def _clean_number(value) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace(" ", "").replace("%", "")
    text = text.replace("٫", ".").replace("٬", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def _row_text(row) -> str:
    return " ".join(_clean_text(value) for value in row if _clean_text(value))


def _find_header_row(sheet, required_groups, max_rows=20):
    # Each group contains alternative labels; at least one label from every group is required.
    for row_no in range(1, min(sheet.max_row, max_rows) + 1):
        values = [_clean_text(sheet.cell(row_no, col).value) for col in range(1, sheet.max_column + 1)]
        if all(any(keyword in value for keyword in group for value in values) for group in required_groups):
            return row_no
    return None


def _find_columns(sheet, header_row, aliases):
    columns = {}
    for col in range(1, sheet.max_column + 1):
        value = _clean_text(sheet.cell(header_row, col).value)
        if not value:
            continue
        normalized_value = _normal_key(value)
        for key, words in aliases.items():
            if key in columns:
                continue
            if any(_normal_key(word) in normalized_value for word in words):
                columns[key] = col
    return columns


def _normal_key(value: str) -> str:
    return (
        value.replace("ي", "ی")
        .replace("ك", "ک")
        .replace("\u200c", " ")
        .replace(" ", "")
        .strip()
        .lower()
    )


_PERSIAN_MONTHS = [
    "فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
    "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند",
]


def _iso_date(value) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean_text(value)
    try:
        return date.fromisoformat(text[:10]).isoformat()
    except ValueError:
        return text


def _report_period(date_value: str) -> str | None:
    text = _clean_text(date_value)
    match = re.match(r"^(\d{4})[-/](\d{1,2})[-/]\d{1,2}$", text)
    if not match:
        return None
    year, month = int(match.group(1)), int(match.group(2))
    if 1300 <= year <= 1500:
        jalali_year, jalali_month = year, month
    else:
        try:
            converted = jdatetime.date.fromgregorian(date=date.fromisoformat(text[:10]))
        except ValueError:
            return None
        jalali_year, jalali_month = converted.year, converted.month
    if not 1 <= jalali_month <= 12:
        return None
    return f"{_PERSIAN_MONTHS[jalali_month - 1]} {jalali_year}"


def _summary_month_columns(sheet):
    normalized_months = [_normal_key(month) for month in _PERSIAN_MONTHS]
    for row_no in range(1, min(sheet.max_row, 12) + 1):
        columns = []
        for col_no in range(1, sheet.max_column + 1):
            label = _normal_key(_clean_text(sheet.cell(row_no, col_no).value))
            if not re.search(r"14\d{2}", label):
                continue
            matching_months = [month for month in normalized_months if month in label]
            if len(matching_months) == 1:
                columns.append((col_no, _clean_text(sheet.cell(row_no, col_no).value)))
        if columns:
            return row_no, columns
    return None, []


def _parse_project_metadata(workbook):
    candidates = ["جلد"]
    sheet = next((workbook[name] for name in candidates if name in workbook.sheetnames), None)
    if sheet is None:
        return {"name": "پروژه مترو", "period_label": "دوره گزارش"}

    contract_sheet = _find_sheet(workbook, ["پیشرفت درآمدی پروژه"])
    contract_value = 0.0
    if contract_sheet is not None:
        contract_value = _clean_number(contract_sheet["D11"].value)

    return {
        "name": _clean_text(sheet["B2"].value) or "پروژه مترو",
        "period_label": _clean_text(sheet["J1"].value) or "دوره گزارش",
        "contract_value": contract_value,
    }


def _parse_summary_metrics(workbook):
    if "خلاصه گزارش" not in workbook.sheetnames:
        raise ValueError("Sheet 'خلاصه گزارش' was not found.")

    sheet = workbook["خلاصه گزارش"]
    summary = {
        "physical_progress": _clean_number(sheet["F5"].value),
        "revenue": _clean_number(sheet["F11"].value),
        "contract_revenue": _clean_number(sheet["F7"].value),
        "adjustment": _clean_number(sheet["F8"].value),
        "adjustment_difference": _clean_number(sheet["F9"].value),
        "other_income": _clean_number(sheet["F10"].value),
        "cost": _clean_number(sheet["F20"].value),
        "revenue_progress": _clean_number(sheet["F12"].value),
        "monthly_finance": [],
    }

    summary["total_revenue"] = summary["revenue"]
    _header_row, month_columns = _summary_month_columns(sheet)
    for col, label in month_columns:
        revenue = _clean_number(sheet.cell(11, col).value)
        cost = _clean_number(sheet.cell(20, col).value)
        if revenue or cost:
            summary["monthly_finance"].append((label, revenue, cost))

    if not summary["monthly_finance"]:
        summary["monthly_finance"] = [
            ("دوره گزارش", summary["total_revenue"], summary["cost"])
        ]

    summary["balance"] = summary["total_revenue"] - summary["cost"]
    return summary


def _find_sheet(workbook, names):
    for name in names:
        if name in workbook.sheetnames:
            return workbook[name]
    return None


def _parse_physical_progress(workbook):
    sheet = _find_sheet(workbook, ["پيشرفت فيزيكي", "پیشرفت فیزیکی", "پیشرفت فیزیکی11"])
    if sheet is None:
        return {}

    header_row = _find_header_row(
        sheet,
        [
            ["فعالیت", "شرح"],
            ["جبهه", "موقعیت"],
            ["واحد"],
        ],
        max_rows=25,
    )
    if header_row is None:
        return {}

    columns = _find_columns(
        sheet,
        header_row,
        {
            "position": ["فعالیت"],
            "zone": ["جبهه های کاری", "جبهه‌هاي كاري", "جبهه کاری"],
            "title": ["شرح فعالیت"],
            "unit": ["واحد"],
            "unit_rate": ["بهای واحد درآمد"],
            "total": ["آخرین برآورد احجام کل", "برآورد احجام کل"],
            "opening": ["مقدار انجام شده تا ابتدای دوره"],
            "cumulative": ["تجمعی"],
            "remaining": ["حجم باقی مانده", "حجم باقیمانده"],
        },
    )

    month_columns = []
    period_header_row = header_row + 1
    first_period_col = columns.get("opening", 0) + 1
    last_period_col = columns.get("cumulative", sheet.max_column + 1)
    if "cumulative" not in columns:
        cumulative_label = _normal_key("تجمعی")
        last_period_col = next(
            (
                col_no
                for col_no in range(first_period_col, sheet.max_column + 1)
                if _normal_key(_clean_text(sheet.cell(period_header_row, col_no).value))
                == cumulative_label
            ),
            last_period_col,
        )
    for col in range(first_period_col, last_period_col):
        label = _clean_text(sheet.cell(period_header_row, col).value)
        normalized_label = _normal_key(label)
        if re.search(r"14\d{2}", normalized_label) and any(
            _normal_key(month) in normalized_label
            for month in ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور", "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
        ):
            month_columns.append((col, label))

    result = {}
    for row_no in range(header_row + 1, sheet.max_row + 1):
        title = _clean_text(sheet.cell(row_no, columns.get("title", 0)).value) if columns.get("title") else ""
        position = _clean_text(sheet.cell(row_no, columns.get("position", 0)).value) if columns.get("position") else ""
        zone = _clean_text(sheet.cell(row_no, columns.get("zone", 0)).value) if columns.get("zone") else ""
        if not title:
            continue

        key = (_normal_key(position), _normal_key(zone), _normal_key(title))
        item = {
            "row_no": row_no,
            "position": position,
            "zone": zone,
            "title": title,
            "unit": _clean_text(sheet.cell(row_no, columns["unit"]).value) if columns.get("unit") else "",
            "unit_rate": _clean_number(sheet.cell(row_no, columns["unit_rate"]).value) if columns.get("unit_rate") else 0,
            "total": _clean_number(sheet.cell(row_no, columns["total"]).value) if columns.get("total") else 0,
            "opening": _clean_number(sheet.cell(row_no, columns["opening"]).value) if columns.get("opening") else 0,
            "remaining": _clean_number(sheet.cell(row_no, columns["remaining"]).value) if columns.get("remaining") else 0,
            "periods": [],
        }

        for col, label in month_columns:
            value = _clean_number(sheet.cell(row_no, col).value)
            if value:
                item["periods"].append((label, value))

        if item["total"] <= 0:
            item["total"] = (
                item["opening"]
                + item["remaining"]
                + sum(value for _period, value in item["periods"])
            )

        result[key] = item

    return result


def _parse_operational_plan(workbook):
    sheet = _find_sheet(workbook, ["برنامه عملیاتی", "برنامه اجرایی"])
    if sheet is None:
        return []

    header_row = _find_header_row(
        sheet,
        [
            ["موقعیت", "جبهه کاری"],
            ["شرح فعالیت", "شرح"],
            ["حجم باقی مانده", "باقی مانده"],
        ],
        max_rows=25,
    )
    if header_row is None:
        return []

    columns = _find_columns(
        sheet,
        header_row,
        {
            "position": ["موقعیت"],
            "zone": ["جبهه کاری", "جبهه"],
            "title": ["شرح فعالیت", "شرح"],
            "remaining": ["حجم باقی مانده", "حجم باقیمانده", "باقی مانده"],
            "start": ["شروع"],
            "finish": ["پایان", "اتمام"],
            "duration": ["مدت"],
            "daily": ["برنامه روزانه", "روزانه", "راندمان"],
        },
    )

    result = []
    fixed_cols = set(columns.values())
    current_position = ""
    current_zone = ""

    # Daily columns are only accepted when a nearby header cell contains a date.
    daily_columns = []
    for col in range(1, sheet.max_column + 1):
        if col in fixed_cols:
            continue
        labels = []
        for r in range(max(1, header_row - 3), header_row + 1):
            value = sheet.cell(r, col).value
            if value is not None:
                labels.append(_clean_text(value))
        date_value = next(
            (sheet.cell(r, col).value for r in range(max(1, header_row - 3), header_row + 1)
             if hasattr(sheet.cell(r, col).value, "strftime")),
            None,
        )
        if date_value is not None:
            daily_columns.append((col, _iso_date(date_value)))
            continue

        label = next((v for v in labels if re.search(r"140[0-9][/\-]", v)), "")
        if label:
            daily_columns.append((col, label))

    for row_no in range(header_row + 1, sheet.max_row + 1):
        title = _clean_text(sheet.cell(row_no, columns.get("title", 0)).value) if columns.get("title") else ""
        raw_zone = _clean_text(sheet.cell(row_no, columns.get("zone", 0)).value) if columns.get("zone") else ""
        raw_position = _clean_text(sheet.cell(row_no, columns.get("position", 0)).value) if columns.get("position") else ""
        if raw_position:
            if _normal_key(raw_position) != _normal_key(current_position):
                current_zone = ""
            current_position = raw_position
        if raw_zone:
            current_zone = raw_zone
        zone = current_zone
        position = current_position
        if not title:
            continue

        item = {
            "row_no": row_no,
            "position": position,
            "zone": zone,
            "title": title,
            "remaining": _clean_number(sheet.cell(row_no, columns["remaining"]).value) if columns.get("remaining") else 0,
            "start": _iso_date(sheet.cell(row_no, columns["start"]).value) if columns.get("start") else "",
            "finish": _iso_date(sheet.cell(row_no, columns["finish"]).value) if columns.get("finish") else "",
            "duration": int(_clean_number(sheet.cell(row_no, columns["duration"]).value)) if columns.get("duration") else 0,
            "daily": _clean_number(sheet.cell(row_no, columns["daily"]).value) if columns.get("daily") else 0,
            "daily_plan": [],
        }

        for col, label in daily_columns:
            value = _clean_number(sheet.cell(row_no, col).value)
            if value:
                item["daily_plan"].append((label, value))

        result.append(item)

    return result


def _parse_detail_cost(workbook):
    sheet = _find_sheet(workbook, ["عملکرد مالی-هزینه‌ها", "عملکرد مالی-هزینه ها"])
    if sheet is None:
        return 0.0

    total_col = next(
        (
            col_no
            for col_no in range(1, sheet.max_column + 1)
            if _normal_key(_clean_text(sheet.cell(3, col_no).value)) in {"مجموع", "جمع"}
        ),
        None,
    )
    if total_col is None:
        return 0.0

    subtotals = []
    for row_no in range(4, sheet.max_row + 1):
        labels = _normal_key(
            " ".join(
                _clean_text(sheet.cell(row_no, col_no).value)
                for col_no in (2, 3)
            )
        )
        if "مجموع" not in labels and "جمع" not in labels:
            continue
        value = _clean_number(sheet.cell(row_no, total_col).value)
        if value:
            subtotals.append(value)
    return subtotals[-1] if subtotals else 0.0


def _parse_financial_entries(workbook):
    sheet = workbook["خلاصه گزارش"] if "خلاصه گزارش" in workbook.sheetnames else None
    if sheet is None:
        return {"revenue": [], "cost": []}

    revenue_entries = []
    cost_entries = []
    _header_row, month_columns = _summary_month_columns(sheet)
    section = None

    for row_no in range(4, sheet.max_row + 1):
        group_label = _normal_key(_clean_text(sheet.cell(row_no, 2).value))
        if "درآمد" in group_label:
            section = "revenue"
        elif "هزینه" in group_label or "هزينه" in group_label:
            section = "cost"
        elif "بالانس" in group_label or "پیشرفت" in group_label:
            section = None

        if section is None:
            continue

        category = _clean_text(sheet.cell(row_no, 3).value)
        normalized_category = _normal_key(category)
        if not category:
            continue
        if normalized_category in {"مجموع", "جمع", "تجمعی", "کل"}:
            section = None
            continue

        entries = revenue_entries if section == "revenue" else cost_entries
        for col_no, period in month_columns:
            amount = _clean_number(sheet.cell(row_no, col_no).value)
            if amount:
                entries.append(
                    {"category": category, "period": period, "amount": amount}
                )

    return {"revenue": revenue_entries, "cost": cost_entries}


def _matches_any(text: str, candidates: list[str]) -> bool:
    normalized_text = _normal_key(text)
    normalized_candidates = {_normal_key(candidate) for candidate in candidates}
    return any(candidate in normalized_text for candidate in normalized_candidates)


def _resource_category_for_sheet(title: str) -> str | None:
    if _matches_any(title, ["ماشین", "ماشين", "دستگاه", "ماشین آلات", "ماشين الات", "ماشینالات"]):
        return "ماشین‌آلات"
    if _matches_any(title, ["نیروی", "نيروي", "انسانی", "انساني", "پرسنل", "شغل"]):
        return "نیروی انسانی"
    if _matches_any(title, ["مواد", "مصالح", "ماده", "مواد و مصالح", "مصالح اصلی"]):
        return "مصالح"
    return None


def _parse_resources(workbook):
    resources = []
    seen = set()

    for sheet in workbook.worksheets:
        category = _resource_category_for_sheet(sheet.title)
        if category is None:
            continue
        _header_row, month_columns = _summary_month_columns(sheet)

        if category in {"ماشین‌آلات", "نیروی انسانی"}:
            title_col = 2
            supply_col = 3
            unit = "دستگاه" if category == "ماشین‌آلات" else "نفر"
            header_row = 3
            for row_no in range(header_row + 1, sheet.max_row + 1):
                title = _clean_text(sheet.cell(row_no, title_col).value)
                if not title:
                    continue
                periods = []
                for col_no, period in month_columns:
                    raw_value = sheet.cell(row_no, col_no).value
                    if raw_value is None:
                        continue
                    periods.append(
                        {
                            "period": period,
                            "required_qty": _clean_number(raw_value),
                            "available_qty": None,
                            "opening_stock": None,
                            "purchase_qty": None,
                            "unit_price": None,
                        }
                    )
                if not periods:
                    continue
                supply_type = _clean_text(sheet.cell(row_no, supply_col).value)
                key = (category, _normal_key(title), _normal_key(supply_type))
                if key in seen:
                    continue
                seen.add(key)
                resources.append(
                    {
                        "category": category,
                        "title": title,
                        "required": max(period["required_qty"] for period in periods),
                        "available": None,
                        "unit": unit,
                        "supply_type": supply_type,
                        "unit_price": None,
                        "periods": periods,
                    }
                )
            continue

        for row_no in range(4, sheet.max_row + 1):
            title = _clean_text(sheet.cell(row_no, 2).value)
            if not title:
                continue
            unit = _clean_text(sheet.cell(row_no, 3).value)
            unit_price_value = sheet.cell(row_no, 4).value
            unit_price = _clean_number(unit_price_value) if unit_price_value not in (None, "-") else None
            periods = []
            for opening_col, period in month_columns:
                opening = sheet.cell(row_no, opening_col).value
                required = sheet.cell(row_no, opening_col + 1).value
                purchase = sheet.cell(row_no, opening_col + 2).value
                if opening is None and required is None and purchase is None:
                    continue
                periods.append(
                    {
                        "period": period,
                        "required_qty": _clean_number(required),
                        "available_qty": _clean_number(opening),
                        "opening_stock": _clean_number(opening),
                        "purchase_qty": _clean_number(purchase),
                        "unit_price": unit_price,
                    }
                )
            if not periods:
                continue
            key = (category, _normal_key(title))
            if key in seen:
                continue
            seen.add(key)
            resources.append(
                {
                    "category": category,
                    "title": title,
                    "required": sum(period["required_qty"] for period in periods),
                    "available": periods[0]["opening_stock"],
                    "unit": unit,
                    "supply_type": "",
                    "unit_price": unit_price,
                    "periods": periods,
                }
            )

    return resources


def _parse_risks(workbook):
    score_header = _normal_key("Risk Score")
    sheet = None
    header_row = None
    for candidate in workbook.worksheets:
        for row_no in range(1, min(candidate.max_row, 20) + 1):
            if any(
                _normal_key(_clean_text(candidate.cell(row_no, col_no).value)) == score_header
                for col_no in range(1, min(candidate.max_column, 30) + 1)
            ):
                sheet = candidate
                header_row = row_no
                break
        if sheet is not None:
            break
    if sheet is None:
        return []

    columns = _find_columns(
        sheet,
        header_row,
        {
            "category": ["گروه ریسک", "دسته ریسک"],
            "title": ["شرح فعالیت", "شرح ریسک"],
            "consequence": ["پیامد ریسک", "پیامد"],
            "existing_controls": ["اقدامات کنترلی موجود", "کنترل موجود"],
            "probability": ["احتمال وقوع", "احتمال"],
            "impact": ["شدت وقوع", "شدت اثر", "شدت"],
            "control": ["کنترل ریسک", "ضریب کنترل"],
            "score": ["Risk Score", "امتیاز ریسک"],
            "action": ["اقدام اصلاحی پیشنهادی", "اقدام اصلاحی"],
        },
    )

    risks = []
    for row_no in range(header_row + 1, sheet.max_row + 1):
        title = _clean_text(sheet.cell(row_no, columns.get("title", 0)).value) if columns.get("title") else ""
        if not title:
            continue

        probability = int(_clean_number(sheet.cell(row_no, columns["probability"]).value)) if columns.get("probability") else 0
        impact = int(_clean_number(sheet.cell(row_no, columns["impact"]).value)) if columns.get("impact") else 0
        control = int(_clean_number(sheet.cell(row_no, columns["control"]).value)) if columns.get("control") else 0
        score_value = _clean_number(sheet.cell(row_no, columns["score"]).value) if columns.get("score") else 0
        risks.append(
            {
                "category": _clean_text(sheet.cell(row_no, columns["category"]).value) if columns.get("category") else "",
                "title": title,
                "consequence": _clean_text(sheet.cell(row_no, columns["consequence"]).value) if columns.get("consequence") else "",
                "existing_controls": _clean_text(sheet.cell(row_no, columns["existing_controls"]).value) if columns.get("existing_controls") else "",
                "probability": probability,
                "impact": impact,
                "control": control,
                "score": int(score_value) if score_value else probability * impact * control,
                "action": _clean_text(sheet.cell(row_no, columns["action"]).value) if columns.get("action") else "",
            }
        )

    return risks


def _upsert_activity(conn, item, physical, row_no):
    key = (
        _normal_key(item["position"]),
        _normal_key(item["zone"]),
        _normal_key(item["title"]),
    )
    p = physical.get(key, {})

    total = p.get("total", 0) or 0
    remaining = p.get("remaining", 0) or 0
    opening = p.get("opening", 0) or 0
    actual_period_qty = sum(value for _period, value in p.get("periods", []))
    actual_total = opening + actual_period_qty
    if total <= 0:
        total = actual_total + remaining if p else 0

    existing = conn.execute(
        """
        SELECT id FROM activity
        WHERE project_id=1 AND source='EXCEL' AND row_no=?
          AND position=? AND zone=? AND title=?
        """,
        (row_no, item["position"], item["zone"], item["title"]),
    ).fetchone()
    if existing is None:
        existing = conn.execute(
            """
            SELECT id FROM activity
            WHERE project_id=1 AND source='EXCEL' AND position=? AND zone=? AND title=?
            ORDER BY row_no LIMIT 1
            """,
            (item["position"], item["zone"], item["title"]),
        ).fetchone()
    if existing is None and row_no is not None:
        manual_override = conn.execute(
            "SELECT id FROM activity WHERE project_id=1 AND row_no=? AND source='MANUAL'",
            (row_no,),
        ).fetchone()
        if manual_override:
            return manual_override["id"]

    manual_actual = 0.0
    if existing:
        manual_actual = float(
            conn.execute(
                "SELECT COALESCE(SUM(quantity),0) FROM activity_daily_actual "
                "WHERE activity_id=? AND source='MANUAL'",
                (existing["id"],),
            ).fetchone()[0]
            or 0
        )
    remaining = max(0.0, float(p.get("remaining", 0) or 0) - manual_actual)

    values = (
        item["position"],
        item["zone"],
        item["title"],
        total,
        item["remaining"],
        remaining,
        p.get("unit", ""),
        item["start"],
        item["finish"],
        item["duration"],
        item["daily"],
        sum(value for _date, value in item["daily_plan"]),
        actual_total + manual_actual,
        actual_total,
        ((actual_total + manual_actual) / total) if total else 0,
        0,
        "NORMAL",
    )

    if existing:
        conn.execute(
            """
            UPDATE activity SET position=?, zone=?, title=?, quantity=?, plan_remaining_qty=?, remaining_qty=?,
            unit=?, start_date=?, finish_date=?, duration_days=?, daily_target=?,
            planned_qty=?, actual_qty=?, baseline_actual_qty=?, progress=?, row_no=?, source='EXCEL'
            WHERE id=?
            """,
            (*values[:15], row_no, existing["id"]),
        )
        activity_id = existing["id"]
    else:
        cursor = conn.execute(
            """
            INSERT INTO activity(
                project_id, row_no, position, zone, title, quantity, plan_remaining_qty, remaining_qty, unit,
                start_date, finish_date, duration_days, daily_target, planned_qty,
                actual_qty, baseline_actual_qty, progress, delay_days, status, source
            ) VALUES(1,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,'NORMAL','EXCEL')
            """,
            (row_no, *values[:16]),
        )
        activity_id = cursor.lastrowid

    conn.execute(
        "DELETE FROM activity_daily_plan WHERE activity_id=? AND source='EXCEL'",
        (activity_id,),
    )
    for plan_date, quantity in item["daily_plan"]:
        conn.execute(
            """
            INSERT INTO activity_daily_plan(activity_id, plan_date, quantity, source)
            VALUES(?,?,?,'EXCEL')
            ON CONFLICT(activity_id, plan_date) DO UPDATE SET
                quantity=excluded.quantity,
                source='EXCEL'
            WHERE activity_daily_plan.source='EXCEL'
            """,
            (activity_id, plan_date, quantity),
        )

    conn.execute(
        "DELETE FROM activity_period WHERE activity_id=? AND source='EXCEL'",
        (activity_id,),
    )
    actual_by_period = {}
    for period, actual in p.get("periods", []):
        actual_by_period[period] = actual_by_period.get(period, 0.0) + actual
    planned_by_period = {}
    for plan_date, planned_value in item["daily_plan"]:
        period = _report_period(plan_date)
        if period:
            planned_by_period[period] = planned_by_period.get(period, 0.0) + planned_value
    all_periods = set(actual_by_period) | set(planned_by_period)
    ordered_periods = sorted(
        all_periods,
        key=lambda period: (
            int(period.split()[-1]),
            _PERSIAN_MONTHS.index(" ".join(period.split()[:-1])),
        ),
    )
    cumulative = opening
    for period in ordered_periods:
        actual = actual_by_period.get(period, 0.0)
        cumulative += actual
        planned = planned_by_period.get(period, 0.0)
        conn.execute(
            """
            INSERT INTO activity_period(
                activity_id, period, planned_qty, actual_qty, cumulative_qty,
                remaining_qty, achievement_pct, source
            ) VALUES(?,?,?,?,?,?,?,'EXCEL')
            ON CONFLICT(activity_id, period) DO UPDATE SET
                planned_qty=excluded.planned_qty,
                actual_qty=excluded.actual_qty,
                cumulative_qty=excluded.cumulative_qty,
                remaining_qty=excluded.remaining_qty,
                achievement_pct=excluded.achievement_pct,
                source='EXCEL'
            WHERE activity_period.source='EXCEL'
            """,
            (
                activity_id,
                period,
                planned,
                actual,
                cumulative,
                max(0.0, total - cumulative),
                (actual / planned * 100) if planned else 0,
            ),
        )

    period_plan_total = sum(planned_by_period.values())
    if period_plan_total <= 0 or not p.get("periods"):
        activity_status = "UNKNOWN"
    else:
        # Use the shared control calculation so retained manual actuals are
        # included when status is recalculated during a workbook re-import.
        achievement = activity_control(conn, activity_id)["achievement_pct"]
        if achievement >= 100:
            activity_status = "NORMAL"
        elif achievement >= 90:
            activity_status = "WARNING"
        else:
            activity_status = "CRITICAL"

    conn.execute(
        "UPDATE activity SET status=? WHERE id=? AND source='EXCEL'",
        (activity_status, activity_id),
    )
    return activity_id


def import_workbook(path, db_path: str | Path | None = None):
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(path)

    workbook_values = openpyxl.load_workbook(path, data_only=True, read_only=False)
    workbook_formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)

    metadata = _parse_project_metadata(workbook_values)
    metrics = _parse_summary_metrics(workbook_values)
    physical = _parse_physical_progress(workbook_values)
    operational = _parse_operational_plan(workbook_values)
    detail_cost = _parse_detail_cost(workbook_values)
    resources = _parse_resources(workbook_values)
    risks = _parse_risks(workbook_values)
    financial_entries = _parse_financial_entries(workbook_values)
    formula_errors = scan_formula_errors(workbook_formulas, workbook_values)

    conn = connect(db_path)
    try:
        conn.execute("BEGIN")

        conn.execute(
            """
            INSERT INTO project(id, name, code, contract_value, period_label, status)
            VALUES (1, ?, 'MB-1405-01', ?, ?, 'ACTIVE')
            ON CONFLICT(id) DO UPDATE SET
                name=excluded.name,
                contract_value=excluded.contract_value,
                period_label=excluded.period_label,
                status=excluded.status
            """,
            (metadata["name"], metadata["contract_value"], metadata["period_label"]),
        )

        conn.execute("DELETE FROM kpi WHERE project_id=1")
        conn.execute(
            """
            INSERT INTO kpi(project_id, physical_progress, revenue, cost, balance, revenue_progress)
            VALUES(1,?,?,?,?,?)
            """,
            (
                metrics["physical_progress"],
                metrics["total_revenue"],
                metrics["cost"],
                metrics["balance"],
                metrics["revenue_progress"],
            ),
        )

        imported_months = []
        for month_name, revenue, cost in metrics["monthly_finance"]:
            imported_months.append(month_name)
            existing_month = conn.execute(
                "SELECT id FROM monthly_finance WHERE project_id=1 AND month=? AND source='EXCEL'",
                (month_name,),
            ).fetchone()
            if existing_month:
                conn.execute(
                    "UPDATE monthly_finance SET revenue=?,cost=? WHERE id=?",
                    (revenue, cost, existing_month["id"]),
                )
            else:
                conn.execute(
                    "INSERT INTO monthly_finance(project_id,month,revenue,cost,source) VALUES(1,?,?,?,'EXCEL')",
                    (month_name, revenue, cost),
                )
        if imported_months:
            placeholders = ",".join("?" for _ in imported_months)
            conn.execute(
                f"DELETE FROM monthly_finance WHERE project_id=1 AND source='EXCEL' AND month NOT IN ({placeholders})",
                imported_months,
            )

        conn.execute("DELETE FROM discrepancy WHERE project_id=1 AND source='EXCEL'")
        conn.execute("DELETE FROM revenue_entry WHERE project_id=1 AND source='EXCEL'")
        conn.execute("DELETE FROM cost_entry WHERE project_id=1 AND source='EXCEL'")
        conn.execute("DELETE FROM activity_dependency WHERE project_id=1 AND source='INFERRED'")
        conn.execute("DELETE FROM physical_progress_entry WHERE project_id=1 AND source='EXCEL'")

        imported_resource_ids = set()
        for item in resources:
            existing_resource = conn.execute(
                """
                SELECT id FROM resource
                WHERE project_id=1 AND category=? AND title=?
                  AND COALESCE(supply_type,'')=? AND source='EXCEL'
                """,
                (item["category"], item["title"], item["supply_type"]),
            ).fetchone()
            resource_values = (
                item["required"], item["available"], item["unit"],
                item["supply_type"], item["unit_price"],
            )
            if existing_resource:
                resource_id = existing_resource["id"]
                conn.execute(
                    """
                    UPDATE resource SET required=?,available=?,unit=?,supply_type=?,unit_price=?
                    WHERE id=?
                    """,
                    (*resource_values, resource_id),
                )
            else:
                cursor = conn.execute(
                    """
                    INSERT INTO resource(
                        project_id,category,title,required,available,unit,supply_type,unit_price,source
                    ) VALUES(1,?,?,?,?,?,?,?,'EXCEL')
                    """,
                    (item["category"], item["title"], *resource_values),
                )
                resource_id = cursor.lastrowid

            imported_resource_ids.add(resource_id)
            imported_periods = []
            for period in item["periods"]:
                imported_periods.append(period["period"])
                conn.execute(
                    """
                    INSERT INTO resource_period(
                        resource_id, period, required_qty, available_qty,
                        opening_stock, purchase_qty, unit_price, source
                    ) VALUES(?,?,?,?,?,?,?,'EXCEL')
                    ON CONFLICT(resource_id,period,source) DO UPDATE SET
                        required_qty=excluded.required_qty,
                        available_qty=excluded.available_qty,
                        opening_stock=excluded.opening_stock,
                        purchase_qty=excluded.purchase_qty,
                        unit_price=excluded.unit_price
                    """,
                    (
                        resource_id, period["period"], period["required_qty"],
                        period["available_qty"], period["opening_stock"],
                        period["purchase_qty"], period["unit_price"],
                    ),
                )
            if imported_periods:
                placeholders = ",".join("?" for _ in imported_periods)
                conn.execute(
                    f"DELETE FROM resource_period WHERE resource_id=? AND source='EXCEL' AND period NOT IN ({placeholders})",
                    (resource_id, *imported_periods),
                )


        # Remove Excel-owned resources that disappeared from a successfully parsed
        # resource section. Never clean the table when parsing returned no rows.
        if resources and imported_resource_ids:
            placeholders = ",".join("?" for _ in imported_resource_ids)
            conn.execute(
                f"DELETE FROM resource WHERE project_id=1 AND source='EXCEL' AND id NOT IN ({placeholders})",
                tuple(imported_resource_ids),
            )

        imported_risk_ids = set()

        for risk in risks:
            risk_values = (
                risk["consequence"], risk["existing_controls"], risk["probability"],
                risk["impact"], risk["control"], risk["score"], risk["action"],
            )
            existing_risk = conn.execute(
                "SELECT id FROM risk WHERE project_id=1 AND category=? AND title=? AND source='EXCEL'",
                (risk["category"], risk["title"]),
            ).fetchone()
            if existing_risk:
                imported_risk_ids.add(existing_risk["id"])
                conn.execute(
                    """
                    UPDATE risk SET consequence=?,existing_controls=?,probability=?,impact=?,
                        control=?,score=?,action=? WHERE id=?
                    """,
                    (*risk_values, existing_risk["id"]),
                )
            else:
                cursor = conn.execute(
                    """
                    INSERT INTO risk(
                        project_id,category,title,consequence,existing_controls,
                        probability,impact,control,score,action,source
                    ) VALUES(1,?,?,?,?,?,?,?,?,?,'EXCEL')
                    """,
                    (risk["category"], risk["title"], *risk_values),
                )
                imported_risk_ids.add(cursor.lastrowid)

        # Preserve manually maintained risks and avoid deleting everything when
        # the workbook's risk sheet could not be parsed.
        if risks and imported_risk_ids:
            placeholders = ",".join("?" for _ in imported_risk_ids)
            conn.execute(
                f"DELETE FROM risk WHERE project_id=1 AND source='EXCEL' AND id NOT IN ({placeholders})",
                tuple(imported_risk_ids),
            )

        for entry in financial_entries["revenue"]:
            conn.execute(
                "INSERT INTO revenue_entry(project_id, activity_id, period, category, amount, source) VALUES(1, NULL, ?, ?, ?, 'EXCEL')",
                (entry["period"], entry["category"], entry["amount"]),
            )

        for entry in financial_entries["cost"]:
            conn.execute(
                "INSERT INTO cost_entry(project_id, activity_id, period, category, amount, source) VALUES(1, NULL, ?, ?, ?, 'EXCEL')",
                (entry["period"], entry["category"], entry["amount"]),
            )

        for item in physical.values():
            periods = item["periods"] or [("دوره گزارش", 0.0)]
            for period, actual in periods:
                conn.execute(
                    """
                    INSERT INTO physical_progress_entry(
                        project_id, source_row_no, position, zone, work_package, unit, unit_rate, total_qty,
                        opening_qty, actual_qty, remaining_qty, period, source
                    ) VALUES(1,?,?,?,?,?,?,?,?,?,?,?,'EXCEL')
                    """,
                    (
                        item["row_no"], item["position"], item["zone"], item["title"],
                        item["unit"], item["unit_rate"], item["total"],
                        item["opening"], actual, item["remaining"], period,
                    ),
                )

        # Do not destroy the existing model when a workbook layout cannot be parsed.
        if operational:
            imported_activity_ids = set()
            for item in operational:
                imported_activity_ids.add(
                    _upsert_activity(conn, item, physical, item["row_no"])
                )
            infer_activity_dependencies = __import__('app.services.control', fromlist=['infer_activity_dependencies']).infer_activity_dependencies
            infer_activity_dependencies(conn, 1)
            schedule_result = calculate_schedule(conn, 1, persist=True)
            if schedule_result["cycle"]:
                conn.execute(
                    "INSERT INTO discrepancy(project_id,title,detail,severity,source) VALUES(1,?,?,?,'EXCEL')",
                    (
                        "چرخه در وابستگی فعالیت‌ها",
                        "محاسبه زمان‌بندی به دلیل چرخه وابستگی انجام نشد: "
                        + "، ".join(schedule_result.get("cycle_titles", [])),
                        "HIGH",
                    ),
                )
            stale_count = conn.execute(
                "SELECT COUNT(*) FROM activity WHERE project_id=1 AND source='EXCEL' "
                "AND id NOT IN ({})".format(
                    ",".join("?" for _ in imported_activity_ids) or "NULL"
                ),
                tuple(imported_activity_ids),
            ).fetchone()[0]
            if stale_count:
                conn.execute(
                    """
                    INSERT INTO discrepancy(project_id,title,detail,severity,source)
                    VALUES(1,?,?,?,'EXCEL')
                    """,
                    (
                        "فعالیت‌های فایل قبلی حفظ شدند",
                        f"{stale_count} فعالیت Excel در فایل جدید یافت نشد؛ برای حفظ تاریخچه حذف نشدند.",
                        "MEDIUM",
                    ),
                )
        else:
            conn.execute(
                    "INSERT INTO discrepancy(project_id,title,detail,severity,source) VALUES(1,?,?,?,'EXCEL')",
                (
                    "ورود برنامه عملیاتی انجام نشد",
                    "ساختار شیت «برنامه عملیاتی» شناسایی نشد؛ داده‌های فعالیت قبلی حفظ شدند.",
                    "HIGH",
                ),
            )

        operational_keys = {
            (
                _normal_key(item["position"]),
                _normal_key(item["zone"]),
                _normal_key(item["title"]),
            )
            for item in operational
        }
        unmatched_physical = len(set(physical) - operational_keys)
        if unmatched_physical:
            conn.execute(
                "INSERT INTO discrepancy(project_id,title,detail,severity,source) VALUES(1,?,?,?,'EXCEL')",
                (
                    "پیشرفت فیزیکی بدون تطبیق یک‌به‌یک",
                    f"{unmatched_physical} بسته از شیت پیشرفت فیزیکی با عنوان فعالیت‌های برنامه عملیاتی تطبیق دقیق ندارند؛ داده خام جداگانه نگهداری شد و به فعالیتی نسبت داده نشد.",
                    "MEDIUM",
                ),
            )

        unmatched_operational = len(operational_keys - set(physical))
        if unmatched_operational:
            conn.execute(
                "INSERT INTO discrepancy(project_id,title,detail,severity,source) VALUES(1,?,?,?,'EXCEL')",
                (
                    "فعالیت‌های برنامه بدون تطبیق پیشرفت فیزیکی",
                    f"{unmatched_operational} فعالیت از برنامه عملیاتی با شیت پیشرفت فیزیکی تطبیق نداشت؛ داده واقعی برای این فعالیت‌ها نامشخص است و وضعیت آن‌ها به‌صورت UNKNOWN نگهداری شد.",
                    "MEDIUM",
                ),
            )

        discrepancy = detect_finance_discrepancy(metrics["cost"], detail_cost)
        if discrepancy:
            conn.execute(
                "INSERT INTO discrepancy(project_id,title,detail,severity,source) VALUES(1,?,?,?,'EXCEL')",
                (discrepancy["title"], discrepancy["detail"], discrepancy["severity"]),
            )

        for error in formula_errors:
            conn.execute(
                """
                INSERT INTO discrepancy(
                    project_id, title, detail, severity, source
                )
                VALUES(1, ?, ?, ?, 'EXCEL')
                """,
                (error["title"], error["detail"], error["severity"]),
            )

        conn.commit()
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()
        workbook_values.close()
        workbook_formulas.close()


if __name__ == "__main__":
    import sys

    if len(sys.argv) != 2:
        raise SystemExit("Usage: python imports/import_excel.py <workbook.xlsx>")

    import_workbook(sys.argv[1])
