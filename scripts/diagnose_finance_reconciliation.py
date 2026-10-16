"""Read-only source-sheet diagnostics for the Excel finance mismatch.

Run from the repository root:
    python scripts/diagnose_finance_reconciliation.py
Or:
    python scripts/diagnose_finance_reconciliation.py path/to/workbook.xlsx

This script does not import data into SQLite or modify the workbook.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
if str(REPOSITORY_ROOT) not in sys.path:
    sys.path.insert(0, str(REPOSITORY_ROOT))

import openpyxl

from imports.import_excel import (
    _clean_number,
    _clean_text,
    _normal_key,
    _parse_detail_cost,
    _summary_month_columns,
)


def find_sheet(workbook, names):
    for name in names:
        if name in workbook.sheetnames:
            return workbook[name]
    return None


def display(value):
    if value is None:
        return ""
    return str(value).replace("\n", " ").strip()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbook",
        nargs="?",
        type=Path,
        default=REPOSITORY_ROOT / "metro-project-control-system.xlsx",
        help="Workbook to inspect (defaults to the project workbook).",
    )
    args = parser.parse_args()
    path = args.workbook.expanduser().resolve()
    if not path.is_file():
        print("ERROR: workbook not found: " + str(path), file=sys.stderr)
        return 2

    values_book = openpyxl.load_workbook(path, data_only=True, read_only=False)
    formulas_book = openpyxl.load_workbook(path, data_only=False, read_only=False)
    try:
        summary = find_sheet(values_book, ["خلاصه گزارش"])
        summary_formulas = (
            formulas_book[summary.title] if summary is not None else None
        )
        detail = find_sheet(
            values_book,
            ["عملکرد مالی-هزینه‌ها", "عملکرد مالی-هزینه ها"],
        )
        detail_formulas = (
            formulas_book[detail.title] if detail is not None else None
        )

        print("WORKBOOK:", path)
        print("SUMMARY SHEET:", summary.title if summary else "NOT FOUND")
        if summary is not None:
            for address, label in [
                ("F20", "summary cost"),
                ("F11", "summary revenue"),
                ("F21", "summary balance"),
            ]:
                cached = summary[address].value
                formula = summary_formulas[address].value if summary_formulas else None
                print(
                    f"SUMMARY {label}: {cached!r}"
                    + (f" | formula={formula!r}" if formula is not None else "")
                )

            header_row, month_columns = _summary_month_columns(summary)
            print("SUMMARY MONTH HEADER:", header_row)
            print("SUMMARY MONTH COLUMNS:", month_columns)
            section = None
            print("\nSUMMARY COST COMPONENTS (rows from the cost section):")
            for row_no in range(4, summary.max_row + 1):
                group = _normal_key(_clean_text(summary.cell(row_no, 2).value))
                if "درآمد" in group:
                    section = "revenue"
                elif "هزینه" in group or "هزينه" in group:
                    section = "cost"
                elif "بالانس" in group or "پیشرفت" in group:
                    section = None
                if section != "cost":
                    continue

                category = _clean_text(summary.cell(row_no, 3).value)
                if not category:
                    continue
                normalized = _normal_key(category)
                if normalized in {"مجموع", "جمع", "تجمعی", "کل"}:
                    section = None
                    continue
                amounts = []
                for col_no, period in month_columns:
                    val = _clean_number(summary.cell(row_no, col_no).value)
                    if val:
                        amounts.append(f"{period}={val:,.2f}")
                if amounts:
                    print(f"  row {row_no}: {category} | " + " | ".join(amounts))

        print("\nDETAIL COST SHEET:", detail.title if detail else "NOT FOUND")
        if detail is None:
            return 1

        total_col = next(
            (
                col_no
                for col_no in range(1, detail.max_column + 1)
                if _normal_key(_clean_text(detail.cell(3, col_no).value))
                in {"مجموع", "جمع"}
            ),
            None,
        )
        print("DETAIL SHEET DIMENSIONS:", detail.max_row, "rows x", detail.max_column, "columns")
        print("DETAIL TOTAL COLUMN:", total_col)
        print("DETAIL PARSER TOTAL (last non-zero sum row):", f"{_parse_detail_cost(values_book):,.2f}")
        if total_col is None:
            print("ERROR: no total column found in row 3.")
            return 1

        print("\nDETAIL COST ROWS (labels in B/C, cached total, total formula):")
        sum_rows = []
        for row_no in range(4, detail.max_row + 1):
            b = _clean_text(detail.cell(row_no, 2).value)
            c = _clean_text(detail.cell(row_no, 3).value)
            cached = detail.cell(row_no, total_col).value
            formula = detail_formulas.cell(row_no, total_col).value if detail_formulas else None
            labels = _normal_key(b + " " + c)
            is_sum = "مجموع" in labels or "جمع" in labels
            if is_sum and _clean_number(cached):
                sum_rows.append((row_no, b, c, _clean_number(cached), formula))
            if not (b or c or cached is not None or formula is not None):
                continue
            if not (b or c) and not is_sum:
                continue
            print(
                f"  row {row_no}: B={b!r}; C={c!r}; "
                f"cached_total={cached!r}; formula={formula!r}"
            )

        print("\nDETAIL SUM/SUBTOTAL CANDIDATES:")
        for row_no, b, c, amount, formula in sum_rows:
            print(
                f"  row {row_no}: {b} | {c} | {amount:,.2f} | formula={formula!r}"
            )
        if sum_rows:
            selected = sum_rows[-1]
            print(
                "\nCURRENT PARSER SELECTS: row "
                + str(selected[0])
                + " = "
                + f"{selected[3]:,.2f}"
            )
        if summary is not None:
            summary_cost = _clean_number(summary["F20"].value)
            detail_cost = _parse_detail_cost(values_book)
            difference = detail_cost - summary_cost
            print("\nRECONCILIATION:")
            print(f"  summary cost F20: {summary_cost:,.2f}")
            print(f"  detail parser total: {detail_cost:,.2f}")
            print(f"  detail minus summary: {difference:,.2f}")
    finally:
        values_book.close()
        formulas_book.close()

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
