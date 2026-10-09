from pathlib import Path
import math
import sqlite3

from openpyxl import load_workbook

from app.db import connect
from app.services.control import activity_control, update_activity_rollup
from imports.import_excel import (
    _parse_detail_cost,
    _parse_physical_progress,
    _parse_resources,
    _report_period,
    import_workbook,
)


def test_workbook_import_includes_resource_sheets(tmp_path):
    workbook = Path(__file__).resolve().parents[1] / "metro-project-control-system.xlsx"
    assert workbook.exists(), "Workbook should exist in project root"
    db_path = tmp_path / "project.db"
    workbook_values = load_workbook(workbook, data_only=True, read_only=True)
    detail_cost = _parse_detail_cost(workbook_values)
    physical_expected = _parse_physical_progress(workbook_values)
    physical_entries_expected = sum(max(1, len(item["periods"])) for item in physical_expected.values())
    resources_expected = _parse_resources(workbook_values)
    resource_periods_expected = sum(len(item["periods"]) for item in resources_expected)
    workbook_values.close()
    workbook_formulas = load_workbook(workbook, data_only=False, read_only=True)
    formula_sheet = workbook_formulas.worksheets[3]
    assert _report_period("2026-09-23") == "مهر 1405"
    assert _report_period("2026-10-23") == "آبان 1405"
    assert formula_sheet["S50"].value == "=I13*S46"
    assert formula_sheet["T50"].value == "=S46*J13"
    assert formula_sheet["S52"].value == "=I15*S46"
    assert formula_sheet["T52"].value == "=S46*J15"
    assert formula_sheet["T58"].value == "=#REF!+J11"
    workbook_formulas.close()

    conn = connect(db_path)
    conn.close()

    import_workbook(str(workbook), db_path=db_path)

    conn = connect(db_path)
    resources = conn.execute("SELECT COUNT(*) FROM resource WHERE project_id=1").fetchone()[0]
    activities = conn.execute("SELECT COUNT(*) FROM activity WHERE project_id=1").fetchone()[0]
    risks = conn.execute("SELECT COUNT(*) FROM risk WHERE project_id=1").fetchone()[0]
    resource_period_count = conn.execute(
        "SELECT COUNT(*) FROM resource_period WHERE source='EXCEL'"
    ).fetchone()[0]
    false_resource_headers = conn.execute(
        "SELECT COUNT(*) FROM resource WHERE title IN ('نوع دستگاه','سمت')"
    ).fetchone()[0]
    cement = conn.execute(
        "SELECT unit,unit_price,required,available FROM resource WHERE project_id=1 AND title='سیمان'"
    ).fetchone()
    cement_periods = {
        row["period"]: (row["opening_stock"], row["required_qty"], row["purchase_qty"])
        for row in conn.execute(
            """
            SELECT rp.period,rp.opening_stock,rp.required_qty,rp.purchase_qty
            FROM resource_period rp JOIN resource r ON r.id=rp.resource_id
            WHERE r.title='سیمان' AND rp.source='EXCEL'
            """
        )
    }
    mixer = conn.execute(
        "SELECT unit,supply_type FROM resource WHERE project_id=1 AND title='تراک میکسر' AND supply_type='استیجاری'"
    ).fetchone()
    physical_entries = conn.execute(
        "SELECT COUNT(*) FROM physical_progress_entry WHERE project_id=1"
    ).fetchone()[0]
    physical_mapping_warnings = conn.execute(
        "SELECT COUNT(*) FROM discrepancy WHERE project_id=1 AND title='پیشرفت فیزیکی بدون تطبیق یک‌به‌یک'"
    ).fetchone()[0]
    formula_errors = conn.execute(
        "SELECT COUNT(*) FROM discrepancy WHERE project_id=1 AND title='خطای فرمول Excel'"
    ).fetchone()[0]
    formula_error_details = [
        row[0]
        for row in conn.execute(
            "SELECT detail FROM discrepancy WHERE project_id=1 AND title='خطای فرمول Excel'"
        )
    ]
    risk_score_mismatches = conn.execute(
        "SELECT COUNT(*) FROM risk WHERE project_id=1 AND score<>probability*impact*control"
    ).fetchone()[0]
    database_integrity = conn.execute("PRAGMA integrity_check").fetchone()[0]
    foreign_key_errors = conn.execute("PRAGMA foreign_key_check").fetchall()
    first_risk = conn.execute(
        "SELECT category,probability,impact,control,score,consequence,existing_controls "
        "FROM risk WHERE project_id=1 ORDER BY id LIMIT 1"
    ).fetchone()
    manpower = conn.execute("SELECT COUNT(*) FROM resource WHERE project_id=1 AND category='نیروی انسانی'").fetchone()[0]
    tunnel4 = conn.execute(
        "SELECT COUNT(*) FROM activity WHERE project_id=1 AND (position LIKE '%تونل%' OR zone LIKE '%تونل%' OR title LIKE '%تونل%') AND (position LIKE '%4%' OR zone LIKE '%4%' OR title LIKE '%4%')"
    ).fetchone()[0]
    tunnel4_transfer = conn.execute(
        "SELECT title,start_date,finish_date FROM activity WHERE project_id=1 AND row_no=13"
    ).fetchone()
    revenue_entries = conn.execute("SELECT COUNT(*) FROM revenue_entry WHERE project_id=1").fetchone()[0]
    cost_entries = conn.execute("SELECT COUNT(*) FROM cost_entry WHERE project_id=1").fetchone()[0]
    dependencies = conn.execute("SELECT COUNT(*) FROM activity_dependency WHERE project_id=1").fetchone()[0]
    monthly = {
        row["month"]: (row["revenue"], row["cost"])
        for row in conn.execute("SELECT month,revenue,cost FROM monthly_finance WHERE project_id=1")
    }
    financial_totals = conn.execute(
        "SELECT revenue,cost,balance,revenue_progress FROM kpi WHERE project_id=1"
    ).fetchone()
    contract_value = conn.execute(
        "SELECT contract_value FROM project WHERE id=1"
    ).fetchone()[0]
    revenue_by_period = dict(
        conn.execute(
            "SELECT period,SUM(amount) FROM revenue_entry WHERE project_id=1 GROUP BY period"
        ).fetchall()
    )
    cost_by_period = dict(
        conn.execute(
            "SELECT period,SUM(amount) FROM cost_entry WHERE project_id=1 GROUP BY period"
        ).fetchall()
    )
    balance_category_entries = conn.execute(
        "SELECT COUNT(*) FROM cost_entry WHERE project_id=1 AND category LIKE '%بالانس%'"
    ).fetchone()[0]
    work_groups = {
        row["row_no"]: (row["position"], row["zone"])
        for row in conn.execute(
            "SELECT row_no,position,zone FROM activity WHERE project_id=1 AND row_no BETWEEN 4 AND 14"
        )
    }
    invalid_dependencies = conn.execute(
        """
        SELECT COUNT(*)
        FROM activity_dependency d
        JOIN activity child ON child.id=d.activity_id
        JOIN activity parent ON parent.id=d.predecessor_activity_id
        WHERE child.position IS NULL OR child.position='' OR child.zone IS NULL OR child.zone=''
           OR parent.position IS NULL OR parent.position='' OR parent.zone IS NULL OR parent.zone=''
           OR child.position<>parent.position OR child.zone<>parent.zone
        """
    ).fetchone()[0]
    physical_row = conn.execute(
        """
        SELECT position,zone,work_package,unit_rate,total_qty,opening_qty,
               actual_qty,remaining_qty
        FROM physical_progress_entry
        WHERE source_row_no=10 AND period='مهر 1405'
        """
    ).fetchone()
    activity_row = conn.execute(
        """
        SELECT quantity,plan_remaining_qty,remaining_qty,actual_qty,baseline_actual_qty,progress
        FROM activity
        WHERE position='تونل شماره 3' AND zone='خط غربی' AND title='بتن ریزی زیر داکت'
        """
    ).fetchone()
    planned_periods = {
        row["period"]: (row["planned_qty"], row["actual_qty"], row["achievement_pct"])
        for row in conn.execute(
            """
            SELECT period,planned_qty,actual_qty,achievement_pct FROM activity_period
            WHERE activity_id=(SELECT id FROM activity WHERE project_id=1 AND row_no=4)
              AND source='EXCEL'
            """
        )
    }
    imported_activity_status = conn.execute(
        "SELECT status FROM activity WHERE project_id=1 AND row_no=4"
    ).fetchone()[0]
    conn.close()

    assert activities >= 10, "At least the major operational activities should be imported"
    assert resources >= 10, "Machine, manpower, and material sheets should be imported into resources"
    assert manpower > 0, "Manpower resource sheet should be imported"
    assert resource_period_count == resource_periods_expected > 0
    assert false_resource_headers == 0
    assert tuple(cement[:2]) == ("تن", 28_000_000)
    assert cement_periods == {
        "مهر 1405": (0.0, 836.0168272000001, 836.0168272000001),
        "آبان 1405": (0.0, 525.20614576, 525.20614576),
    }
    assert tuple(mixer) == ("دستگاه", "استیجاری")
    assert risks == 7, "Every project risk from the workbook should be imported"
    assert physical_entries == physical_entries_expected > 0
    assert physical_row is not None
    assert tuple(physical_row[:4]) == ("تونل شماره 3", "خط غربی", "بتن ریزی زیر داکت", 13_314_220.981009783)
    assert math.isclose(physical_row[4], 969, abs_tol=1e-9)
    assert math.isclose(physical_row[5], 0, abs_tol=1e-9)
    assert math.isclose(physical_row[6], 836, abs_tol=1e-9)
    assert math.isclose(physical_row[7], 0, abs_tol=1e-9)
    assert activity_row is not None
    assert math.isclose(activity_row[0], 969, abs_tol=1e-9)
    assert math.isclose(activity_row[3], 969, abs_tol=1e-9)
    assert math.isclose(activity_row[4], 969, abs_tol=1e-9)
    assert physical_mapping_warnings > 0, "Unmatched physical progress must be visible rather than silently discarded"
    assert formula_errors == 1, "Only the one unresolved source formula should remain visible"
    assert len(formula_error_details) == 1 and "T58" in formula_error_details[0]
    assert tuple(first_risk[:5]) == ("مالی", 5, 5, 5, 125)
    assert first_risk["consequence"] and first_risk["existing_controls"] == "--"
    assert tunnel4 >= 1, "Tunnel 4 operational activity should be imported"
    assert tuple(tunnel4_transfer) == ("انتقال ریل به دهانه تونل", "2026-09-29", "2026-10-18")
    assert revenue_entries == 6, "Revenue line items from both months should be imported"
    assert cost_entries == 12, "All six cost categories from both months should be imported"
    assert balance_category_entries == 0, "Project balance must not be stored as a cost"
    assert dependencies > 0, "Activity dependencies should be inferred from workbook sequencing"
    assert set(monthly) == {"مهر 1405", "آبان 1405"}
    assert work_groups[5] == ("محدوده روباز", "خط غربی")
    assert work_groups[7] == ("تونل شماره 2", "خط شرقی")
    assert work_groups[12] == ("تونل شماره 3", "خط غربی")
    assert invalid_dependencies == 0, "Dependencies must not cross or infer blank work groups"
    assert math.isclose(sum(value[0] for value in monthly.values()), financial_totals["revenue"], abs_tol=0.1)
    assert math.isclose(sum(value[1] for value in monthly.values()), financial_totals["cost"], abs_tol=0.1)
    for period, (revenue, cost) in monthly.items():
        assert math.isclose(revenue_by_period[period], revenue, abs_tol=0.1)
        assert math.isclose(cost_by_period[period], cost, abs_tol=0.1)
    assert math.isclose(financial_totals["balance"], financial_totals["revenue"] - financial_totals["cost"], abs_tol=0.1)
    assert math.isclose(financial_totals["revenue_progress"], 0.14962598867734503, abs_tol=1e-12)
    assert math.isclose(contract_value, 2_856_172_028_562, abs_tol=0.1)
    assert planned_periods["مهر 1405"] == (864, 864, 100)
    assert planned_periods["آبان 1405"] == (295, 295, 100)
    assert imported_activity_status == "NORMAL"
    assert math.isclose(detail_cost, 445_741_374_203.4047, abs_tol=0.1)
    assert 12_300_000_000 < detail_cost - financial_totals["cost"] < 12_400_000_000
    assert risk_score_mismatches == 0
    assert database_integrity == "ok"
    assert not foreign_key_errors

    conn = connect(db_path)
    manual_activity_id = conn.execute(
        "SELECT id FROM activity WHERE project_id=1 AND row_no=5"
    ).fetchone()[0]
    conn.execute(
        "INSERT INTO activity_daily_actual(activity_id,actual_date,quantity,source,notes) "
        "VALUES(?,'2027-01-01',5,'MANUAL','re-import regression')",
        (manual_activity_id,),
    )
    conn.execute(
        "INSERT INTO activity_daily_plan(activity_id,plan_date,quantity,source) "
        "VALUES(?,'2027-01-01',7,'MANUAL')",
        (manual_activity_id,),
    )
    conn.execute(
        "INSERT INTO activity_period(activity_id,period,planned_qty,actual_qty,source) "
        "VALUES(?,'manual-period',7,5,'MANUAL')",
        (manual_activity_id,),
    )
    update_activity_rollup(conn, manual_activity_id)
    conn.execute("UPDATE activity SET source='MANUAL' WHERE id=?", (manual_activity_id,))
    conn.execute(
        "INSERT INTO resource(project_id,category,title,required,available,unit,source) "
        "VALUES(1,'مصالح','منبع دستی تست',3,2,'تن','MANUAL')"
    )
    conn.execute(
        "INSERT INTO risk(project_id,category,title,probability,impact,control,score,action,source) "
        "VALUES(1,'تست','ریسک دستی تست',1,2,3,6,'بررسی','MANUAL')"
    )
    conn.execute(
        "INSERT INTO revenue_entry(project_id,period,category,amount,source) "
        "VALUES(1,'manual-period','manual-test',123,'MANUAL')"
    )
    conn.execute(
        "INSERT INTO cost_entry(project_id,period,category,amount,source) "
        "VALUES(1,'manual-period','manual-test',45,'MANUAL')"
    )
    conn.execute(
        "INSERT INTO discrepancy(project_id,title,detail,severity,source) "
        "VALUES(1,'manual-test','keep across re-import','LOW','MANUAL')"
    )
    conn.execute(
        "INSERT INTO monthly_finance(project_id,month,revenue,cost,source) "
        "VALUES(1,'manual-period',123,45,'MANUAL')"
    )
    conn.execute(
        "INSERT INTO activity(project_id,title,quantity,source) "
        "VALUES(1,'فعالیت دستی تست',10,'MANUAL')"
    )
    conn.execute(
        "INSERT INTO activity(project_id,row_no,position,zone,title,quantity,source) "
        "VALUES(1,999,'موقعیت قدیمی','جبهه قدیمی','فعالیت قدیمی فایل',10,'EXCEL')"
    )
    baseline_before = conn.execute(
        "SELECT baseline_actual_qty FROM activity WHERE project_id=1 AND row_no=5"
    ).fetchone()[0]
    conn.commit()
    conn.close()

    import_workbook(str(workbook), db_path=db_path)

    conn = connect(db_path)
    manual_actual = conn.execute(
        "SELECT COUNT(*),SUM(quantity) FROM activity_daily_actual WHERE source='MANUAL'"
    ).fetchone()
    manual_plan = conn.execute(
        "SELECT COUNT(*) FROM activity_daily_plan WHERE source='MANUAL' AND plan_date='2027-01-01'"
    ).fetchone()[0]
    manual_period = conn.execute(
        "SELECT COUNT(*) FROM activity_period WHERE source='MANUAL' AND period='manual-period'"
    ).fetchone()[0]
    manual_activity = conn.execute(
        "SELECT actual_qty,baseline_actual_qty,remaining_qty FROM activity WHERE project_id=1 AND row_no=5"
    ).fetchone()
    workbook_row_count = conn.execute(
        "SELECT COUNT(*) FROM activity WHERE project_id=1 AND row_no=5"
    ).fetchone()[0]
    preserved_manual_rows = {
        "resource": conn.execute("SELECT COUNT(*) FROM resource WHERE source='MANUAL' AND title='منبع دستی تست'").fetchone()[0],
        "risk": conn.execute("SELECT COUNT(*) FROM risk WHERE source='MANUAL' AND title='ریسک دستی تست'").fetchone()[0],
        "revenue": conn.execute("SELECT COUNT(*) FROM revenue_entry WHERE source='MANUAL' AND category='manual-test'").fetchone()[0],
        "cost": conn.execute("SELECT COUNT(*) FROM cost_entry WHERE source='MANUAL' AND category='manual-test'").fetchone()[0],
        "discrepancy": conn.execute("SELECT COUNT(*) FROM discrepancy WHERE source='MANUAL' AND title='manual-test'").fetchone()[0],
        "monthly_finance": conn.execute("SELECT COUNT(*) FROM monthly_finance WHERE source='MANUAL' AND month='manual-period'").fetchone()[0],
        "activity": conn.execute("SELECT COUNT(*) FROM activity WHERE source='MANUAL' AND title='فعالیت دستی تست'").fetchone()[0],
        "stale_activity": conn.execute("SELECT COUNT(*) FROM activity WHERE source='EXCEL' AND row_no=999").fetchone()[0],
        "stale_warning": conn.execute("SELECT COUNT(*) FROM discrepancy WHERE source='EXCEL' AND title='فعالیت‌های فایل قبلی حفظ شدند'").fetchone()[0],
    }
    formula_warning_count = conn.execute(
        "SELECT COUNT(*) FROM discrepancy WHERE project_id=1 AND title='خطای فرمول Excel'"
    ).fetchone()[0]
    formula_warning_sources = [
        row[0]
        for row in conn.execute(
            "SELECT DISTINCT source FROM discrepancy WHERE project_id=1 AND title='خطای فرمول Excel'"
        )
    ]
    conn.close()

    assert tuple(manual_actual) == (1, 5)
    assert manual_plan == 1
    assert manual_period == 1
    assert manual_activity["baseline_actual_qty"] == baseline_before
    assert manual_activity["actual_qty"] == baseline_before + 5
    assert manual_activity["remaining_qty"] == 1361
    assert workbook_row_count == 1
    assert all(count == 1 for count in preserved_manual_rows.values())
    assert formula_warning_count == 1, "Formula warnings must not duplicate on re-import"
    assert formula_warning_sources == ["EXCEL"], "Workbook formula warnings must be tagged as Excel data"


def test_activity_control_separates_period_and_cumulative_actuals(tmp_path):
    conn = connect(tmp_path / "period-control.db")
    conn.execute(
        "INSERT INTO project(id,name) VALUES(1,'Control test')"
    )
    conn.execute(
        """
        INSERT INTO activity(
            project_id,title,quantity,remaining_qty,actual_qty,
            baseline_actual_qty,source
        ) VALUES(1,'Test activity',100,50,50,50,'EXCEL')
        """
    )
    activity_id = conn.execute(
        "SELECT id FROM activity WHERE title='Test activity'"
    ).fetchone()[0]
    conn.executemany(
        """
        INSERT INTO activity_daily_plan(activity_id,plan_date,quantity,source)
        VALUES(?,?,?,'EXCEL')
        """,
        [
            (activity_id, "2026-09-25", 10),
            (activity_id, "2026-11-12", 10),
        ],
    )
    conn.executemany(
        """
        INSERT INTO activity_period(
            activity_id,period,planned_qty,actual_qty,source
        ) VALUES(?,?,?,?, 'EXCEL')
        """,
        [
            (activity_id, "مهر 1405", 10, 8),
            (activity_id, "آبان 1405", 10, 7),
        ],
    )
    conn.executemany(
        """
        INSERT INTO activity_daily_actual(activity_id,actual_date,quantity,source)
        VALUES(?,?,?,'MANUAL')
        """,
        [
            (activity_id, "2026-10-01", 2),
            (activity_id, "2026-12-01", 3),
        ],
    )
    conn.commit()

    control = activity_control(conn, activity_id)
    conn.close()

    assert control["planned_qty"] == 20
    assert control["period_actual_qty"] == 17
    assert control["cumulative_actual_qty"] == 55
    assert control["actual_qty"] == 55
    assert control["variance_qty"] == -3
    assert control["achievement_pct"] == 85
    assert math.isclose(control["physical_progress_pct"], 55, abs_tol=1e-9)
    assert control["remaining_qty"] == 45


def test_physical_progress_migration_allows_duplicate_descriptions_by_position(tmp_path):

    db_path = tmp_path / "legacy.db"
    legacy = sqlite3.connect(db_path)
    legacy.executescript(
        """
        CREATE TABLE project (id INTEGER PRIMARY KEY);
        INSERT INTO project(id) VALUES(1);
        CREATE TABLE activity (
            id INTEGER PRIMARY KEY,
            project_id INTEGER,
            row_no INTEGER,
            title TEXT,
            quantity REAL,
            remaining_qty REAL,
            actual_qty REAL,
            progress REAL,
            status TEXT
        );
        INSERT INTO activity VALUES(1,1,1,'legacy',100,80,25,0.25,'NORMAL');
        CREATE TABLE activity_daily_actual (
            id INTEGER PRIMARY KEY,
            activity_id INTEGER,
            actual_date TEXT,
            quantity REAL,
            source TEXT,
            notes TEXT
        );
        INSERT INTO activity_daily_actual(activity_id,actual_date,quantity,source)
            VALUES(1,'2026-10-09',5,'MANUAL');
        CREATE TABLE physical_progress_entry (
            id INTEGER PRIMARY KEY,
            project_id INTEGER NOT NULL,
            zone TEXT,
            work_package TEXT NOT NULL,
            unit TEXT,
            total_qty REAL DEFAULT 0,
            opening_qty REAL DEFAULT 0,
            actual_qty REAL DEFAULT 0,
            remaining_qty REAL DEFAULT 0,
            period TEXT NOT NULL,
            source TEXT DEFAULT 'EXCEL',
            FOREIGN KEY(project_id) REFERENCES project(id) ON DELETE CASCADE,
            UNIQUE(project_id, zone, work_package, period)
        );
        INSERT INTO physical_progress_entry(project_id,zone,work_package,period)
        VALUES(1,'خط شرقی','مونتاژ خط','مهر 1405');
        """
    )
    legacy.close()

    conn = connect(db_path)
    conn.execute(
        """
        INSERT INTO physical_progress_entry(
            project_id,source_row_no,position,zone,work_package,period
        ) VALUES(1,2,'تونل شماره 2','خط شرقی','مونتاژ خط','مهر 1405')
        """
    )
    assert conn.execute("SELECT COUNT(*) FROM physical_progress_entry").fetchone()[0] == 2
    assert conn.execute("SELECT COUNT(*) FROM physical_progress_entry WHERE source_row_no IS NULL").fetchone()[0] == 1
    legacy_activity = conn.execute(
        "SELECT plan_remaining_qty,remaining_qty,baseline_actual_qty FROM activity WHERE id=1"
    ).fetchone()
    assert tuple(legacy_activity) == (80, 75, 20)
    assert conn.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    conn.close()



def test_reimport_removes_stale_excel_resources_and_risks_but_keeps_manual_rows(tmp_path):
    workbook = Path(__file__).resolve().parents[1] / "metro-project-control-system.xlsx"
    db_path = tmp_path / "stale-records.db"
    import_workbook(str(workbook), db_path=db_path)

    conn = connect(db_path)
    conn.execute(
        "INSERT INTO resource(project_id,category,title,source) VALUES(1,'مصالح','منبع قدیمی Excel','EXCEL')"
    )
    conn.execute(
        "INSERT INTO resource(project_id,category,title,source) VALUES(1,'مصالح','منبع دستی حفظ‌شونده','MANUAL')"
    )
    conn.execute(
        "INSERT INTO risk(project_id,category,title,source) VALUES(1,'آزمایشی','ریسک قدیمی Excel','EXCEL')"
    )
    conn.execute(
        "INSERT INTO risk(project_id,category,title,source) VALUES(1,'آزمایشی','ریسک دستی حفظ‌شونده','MANUAL')"
    )
    conn.commit()
    conn.close()

    import_workbook(str(workbook), db_path=db_path)

    conn = connect(db_path)
    stale_resource = conn.execute(
        "SELECT COUNT(*) FROM resource WHERE title='منبع قدیمی Excel'"
    ).fetchone()[0]
    manual_resource = conn.execute(
        "SELECT COUNT(*) FROM resource WHERE title='منبع دستی حفظ‌شونده' AND source='MANUAL'"
    ).fetchone()[0]
    stale_risk = conn.execute(
        "SELECT COUNT(*) FROM risk WHERE title='ریسک قدیمی Excel'"
    ).fetchone()[0]
    manual_risk = conn.execute(
        "SELECT COUNT(*) FROM risk WHERE title='ریسک دستی حفظ‌شونده' AND source='MANUAL'"
    ).fetchone()[0]
    conn.close()

    assert stale_resource == 0
    assert manual_resource == 1
    assert stale_risk == 0
    assert manual_risk == 1
